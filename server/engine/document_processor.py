import io
import re
import csv
import json
import logging
from typing import Dict, Any, List, Optional
from server.config import settings
from server.providers.groq_llm import GroqLLM

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """
    Ingests and processes business documents (PDF, DOCX, XLSX, CSV, TXT, JSON)
    and extracts structured business facts, policies, pricing, and FAQs.
    """

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        """Extract clean text content based on file extension"""
        ext = filename.lower().split(".")[-1] if "." in filename else "txt"

        try:
            if ext in ["txt", "md", "text"]:
                return file_bytes.decode("utf-8", errors="ignore").strip()

            elif ext == "csv":
                text_stream = io.StringIO(file_bytes.decode("utf-8", errors="ignore"))
                reader = csv.reader(text_stream)
                rows = [" | ".join(row) for row in reader if any(row)]
                return "\n".join(rows).strip()

            elif ext == "json":
                data = json.loads(file_bytes.decode("utf-8", errors="ignore"))
                return json.dumps(data, indent=2)

            elif ext == "pdf":
                try:
                    import pypdf
                    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                    pages = [page.extract_text() or "" for page in reader.pages]
                    return "\n\n".join(pages).strip()
                except Exception as e:
                    logger.warning(f"pypdf extraction failed for {filename}: {e}")
                    return file_bytes.decode("utf-8", errors="ignore")[:5000]

            elif ext in ["docx", "doc"]:
                try:
                    import docx
                    doc = docx.Document(io.BytesIO(file_bytes))
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    return "\n".join(paragraphs).strip()
                except Exception as e:
                    logger.warning(f"docx extraction failed for {filename}: {e}")
                    return file_bytes.decode("utf-8", errors="ignore")[:5000]

            elif ext in ["xlsx", "xls"]:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                    sheet_texts = []
                    for sheet in wb.worksheets:
                        rows = []
                        for row in sheet.iter_rows(values_only=True):
                            vals = [str(v) for v in row if v is not None]
                            if vals:
                                rows.append(" | ".join(vals))
                        if rows:
                            sheet_texts.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
                    return "\n\n".join(sheet_texts).strip()
                except Exception as e:
                    logger.warning(f"openpyxl extraction failed for {filename}: {e}")
                    return ""

            else:
                return file_bytes.decode("utf-8", errors="ignore").strip()

        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {e}")
            return ""

    @staticmethod
    def extract_structured_knowledge(doc_text: str, filename: str) -> Dict[str, Any]:
        """
        Uses LLM to analyze the document text and extract:
        - Document summary & category
        - Business facts & policies
        - Verified FAQs
        """
        if not doc_text.strip():
            return {
                "category": "General",
                "summary": "Empty or unparseable document",
                "facts": [],
                "faqs": []
            }

        sample_text = doc_text[:4000]

        llm = GroqLLM()
        extraction_prompt = f"""You are an expert Document Intelligence Analyst for Saadhyam AI.
Analyze the following document text and extract structured business facts and FAQs.
Filename: {filename}

Document Content:
\"\"\"
{sample_text}
\"\"\"

Output strictly a JSON object with this exact structure:
{{
  "category": "Pricing & Offerings | Company Policy / SOP | Product Catalog | FAQ Document | General Information",
  "summary": "1-2 sentence overview of the document",
  "facts": ["Key verified business fact 1", "Key verified business fact 2"],
  "faqs": [
    {{"question": "Customer question", "answer": "Exact verified answer from document", "category": "Pricing/General/Booking"}}
  ]
}}
Do NOT output markdown code blocks (```json) or thinking tags. Output valid JSON only."""

        try:
            res = llm.client.chat.completions.create(
                model=llm.model,
                messages=[
                    {"role": "system", "content": "You are a precise JSON extractor. Output valid JSON only without markdown or think tags."},
                    {"role": "user", "content": extraction_prompt}
                ],
                temperature=0.1,
                max_tokens=600
            )
            raw = res.choices[0].message.content.strip()
            if "<think>" in raw and "</think>" in raw:
                raw = raw.split("</think>")[-1].strip()
            if raw.startswith("```"):
                raw = re.sub(r"^```(?:json)?\s*", "", raw)
                raw = re.sub(r"\s*```$", "", raw)

            return json.loads(raw)
        except Exception as e:
            logger.warning(f"LLM structured knowledge extraction notice: {e}")
            lines = [l.strip() for l in doc_text.split("\n") if len(l.strip()) > 15]
            return {
                "category": "General Information",
                "summary": f"Uploaded document: {filename}",
                "facts": lines[:5],
                "faqs": []
            }
